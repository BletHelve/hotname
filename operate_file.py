# This Python file uses the following encoding: utf-8

from threading import Thread
from win32com import client as wc
import docx
import pythoncom
import os
import sys
import xlrd

LACK = -1
FAIL = 0
SUCCESS = 1
REPEAT = 2

ID_INDEX = 0
NAME_INDEX = 1

# 打包之后当前路径
if getattr(sys, 'frozen', None):
    basedir = sys._MEIPASS
else:
    basedir = os.path.dirname(__file__)
#


def get_tb(excel_path):  # 取得excel里的表格
    table = []
    excel_path = os.path.join(basedir, excel_path)
    sheet = xlrd.open_workbook(excel_path).sheet_by_index(0)
    row_cnt = sheet.nrows
    col_names = []
    for name_cell in sheet.row(0):
        col_names.append(name_cell.value)  # 放入列名（属性）
    for row in range(1, row_cnt):
        stu_row = []
        for col in range(len(col_names)):
            cell_type = sheet.cell(row, col).ctype
            val = sheet.cell_value(row, col)
            if cell_type == 2:  # type is number
                val = str(int(val))
            stu_row.append(val)
        table.append(stu_row)
    return table, col_names  # 返回列名，为了适应不同的表，目前可以忽略


tb, col_names = get_tb('table.xlsx')


class OperationFile(Thread):
    __file_list, __string_list = [], []
    __prev_dir, __filename = '', ''
    __state = 0

    def __init__(self, prev_dir, filename, string_list):
        Thread.__init__(self)
        self.__filename = filename
        self.__prev_dir = prev_dir
        self.__string_list = string_list

    def get_file_list(self):
        return self.__file_list

    def operation(self, operation):
        return {'remove': self.remove,
                'rename': self.rename,
                'alt': self.alt
                }[operation]()

    def remove(self):
        filename = os.path.splitext(self.__filename)[0]  # 去后缀名
        if filename in self.__file_list:
            self.__file_list.remove(filename)
            return SUCCESS
        else:
            return FAIL

    def rename(self):
        print(self.__prev_dir + '\\' + self.__filename, end=' ==> ')
        print(self.name_match())
        if self.__state == SUCCESS:
            print('文件重命名成功')
        else:
            print('文件重命名失败:')
            if self.__state == LACK:
                print('文件提供的信息，不足以进行重命名')
            elif self.__state == REPEAT:
                print('将要修改的名称与此文件夹中其他文件名重复')
        return self.__state

    def alt(self):  # 直接根据文件内容进行重命名
        i = self.match_file(self.__prev_dir + '\\' + self.__filename)
        if i != -1:
            print(self.__os_rename(self.__filename, tb[i][ID_INDEX], tb[i][NAME_INDEX]))
            return SUCCESS
        else:
            print('没有改变关键信息，不需要重命名')
            return FAIL

    def name_match(self):
        i = self.match_str(self.__filename)  # 文件名与个人信息表匹配
        if i == -1:  # 没有在文件名中找到姓名和学号，进入word里直接查询
            i = self.match_file(self.__prev_dir + '\\' + self.__filename)
        if i != -1:
            return self.__os_rename(self.__filename, tb[i][ID_INDEX], tb[i][NAME_INDEX])
        else:
            self.__state = LACK

    def __os_rename(self, filename, num, name):
        new_name = num+name  # 新名称
        for s in self.__string_list:
            new_name = new_name+s
        if new_name not in self.__file_list:  # 文件名没有重复
            self.__file_list.append(new_name)
            extension = os.path.splitext(filename)[1]  # 文件后缀名
            os.renames(self.__prev_dir + '\\' + filename,
                       self.__prev_dir + '\\' + new_name + extension)
            self.__state = SUCCESS
        else:
            self.__state = REPEAT
        return new_name

    def match_file(self, file_path):  # 找到文件中的姓名，学号
        extension = os.path.splitext(file_path)[1]  # 获取后缀名
        return {'.doc': self.__match_doc,  # 返回表格的行数
                '.docx': self.__match_docx,
                }[extension](file_path)

    def __match_doc(self, doc_path):
        # todo 这个是io操作，在这里实现切换
        pythoncom.CoInitialize()
        word = wc.Dispatch('Word.Application')
        doc = word.Documents.Open(doc_path)  # 目标路径下的文件
        reg_dir = basedir+'\\reg\\' + os.path.basename(doc_path) + 'x'
        doc.SaveAs(reg_dir, 16)  # 转化成docx，放在reg文件下
        i = self.__match_docx(reg_dir)
        doc.Close()
        word.Quit()
        os.remove(reg_dir)
        return i

    def __match_docx(self, dox_path):
        document = docx.Document(dox_path)
        for paragraph in document.paragraphs:
            line = str(paragraph.text).strip(' ')  # 去掉空格
            i = self.match_str(line)
            if i != -1:
                return i

    @staticmethod
    def match_str(string):  # 名称匹配
        for i in range(len(tb)):
            if tb[i][NAME_INDEX] in string:
                return i
            if tb[i][ID_INDEX] in string:
                return i
        return -1




